"""HERD-9: import the herd from the milking-parlor software's
"ALL HERD" Excel export.

Layout of the export: a title row, a header row whose first cell is
"Num" and fifth is "Status", a second header row (Year / Date / …),
then one row per cow:

    Num | Birth | Group | AGE | Status | Laction | Last Calve | DIM |
    Breeding | Days | Avg | AGE | Age

Only Num, Birth, Group, Status, Laction, Last Calve and Breeding are
read — the rest (ages, DIM, days pregnant, milk average) are derived
again by the app.

The parlor's group numbers and status codes are mapped to ours on a
preview screen (defaults in DEFAULT_STATUS_MAP), so nothing about the
client's setup is hard-coded here.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime

from openpyxl import load_workbook

from app.extensions import db
from app.models.herd import (
    BreedingEvent, BreedingStatusChange, Cow, CowMovement,
)

# Parlor status code (lower-case) → our status code. "" = no status.
DEFAULT_STATUS_MAP = {
    "calf": Cow.BS_CALF,
    "b-clf": Cow.BS_B_CALF,
    "b-calf": Cow.BS_B_CALF,
    "heifr": Cow.BS_HEIFER,
    "heifer": Cow.BS_HEIFER,
    "insem": Cow.BS_INSEM,
    "preg": Cow.BS_PREG,
    "preg2": Cow.BS_PREG2,
    "dry": Cow.BS_DRY,
    "abort": Cow.BS_ABORT,
    "open": Cow.BS_OPEN,
    "fresh": Cow.BS_MILKING,
    "early": Cow.BS_MILKING,
    "ready": Cow.BS_MILKING,
    "lead": Cow.BS_WAITING,
    "none": "",
    "": "",
}

# Column positions in the export.
COL_NUM, COL_BIRTH, COL_GROUP, COL_STATUS, COL_LACT, COL_CALVE, COL_INSEM = 0, 1, 2, 4, 5, 6, 8


class ImportError_(ValueError):
    """Bad file — message is Arabic and safe to flash."""


@dataclass
class ParseResult:
    rows: list[dict] = field(default_factory=list)          # to import
    duplicates: list[str] = field(default_factory=list)     # repeated in the file
    existing: list[str] = field(default_factory=list)       # already in the system
    total: int = 0                                           # rows with a tag

    @property
    def groups(self) -> list[str]:
        return sorted({r["group"] for r in self.rows},
                      key=lambda g: (not g.isdigit(), int(g) if g.isdigit() else 0, g))

    @property
    def statuses(self) -> list[str]:
        return sorted({r["status"] for r in self.rows})

    def to_json(self) -> dict:
        return {"rows": self.rows, "duplicates": self.duplicates,
                "existing": self.existing, "total": self.total}

    @classmethod
    def from_json(cls, data: dict) -> "ParseResult":
        return cls(rows=data["rows"], duplicates=data["duplicates"],
                   existing=data["existing"], total=data["total"])


def _text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _date(v) -> str | None:
    """ISO date string, or None if empty / unreadable."""
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def _int(v) -> int:
    try:
        return max(0, int(float(v)))
    except (TypeError, ValueError):
        return 0


def check_upload(file_storage) -> None:
    name = (file_storage.filename or "").lower()
    if not name.endswith(".xlsx"):
        raise ImportError_("الملف لازم يكون إكسيل بامتداد .xlsx.")
    head = file_storage.stream.read(4)
    file_storage.stream.seek(0)
    if head != b"PK\x03\x04":
        raise ImportError_("الملف مش ملف إكسيل صالح.")


def parse_workbook(stream) -> ParseResult:
    try:
        wb = load_workbook(stream, read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 — any unreadable file
        raise ImportError_("مش قادر أقرا الملف — اتأكد إنه تصدير القطيع من برنامج المحلب.")
    ws = wb.worksheets[0]

    result = ParseResult()
    header_seen = False
    seen = set()
    existing_tags = {t for (t,) in db.session.query(Cow.ear_tag).all()}

    for row in ws.iter_rows(values_only=True):
        row = list(row) + [None] * (13 - len(row))
        if not header_seen:
            if _text(row[COL_NUM]).lower() == "num" and _text(row[COL_STATUS]).lower() == "status":
                header_seen = True
            continue
        tag = _text(row[COL_NUM])
        if not tag or not any(ch.isdigit() for ch in tag):
            continue  # second header row, blanks
        if not (_text(row[COL_GROUP]) or _text(row[COL_STATUS]) or row[COL_BIRTH]):
            continue  # the export's total line ("T  754") — not a cow
        result.total += 1
        if tag in seen:
            result.duplicates.append(tag)
            continue
        seen.add(tag)
        if tag in existing_tags:
            result.existing.append(tag)
            continue
        result.rows.append({
            "tag": tag,
            "birth": _date(row[COL_BIRTH]),
            "group": _text(row[COL_GROUP]),
            "status": _text(row[COL_STATUS]),
            "lactation": _int(row[COL_LACT]),
            "last_calving": _date(row[COL_CALVE]),
            "last_insem": _date(row[COL_INSEM]),
        })
    wb.close()

    if not header_seen:
        raise ImportError_("مش لاقي صف العناوين (Num / Status) — ده مش تصدير القطيع المتوقع.")
    return result


def default_status_for(parlor_status: str) -> str:
    return DEFAULT_STATUS_MAP.get(parlor_status.strip().lower(), "")


def import_rows(rows: list[dict], group_map: dict[str, int],
                status_map: dict[str, str], user_id: int | None) -> int:
    """Create the cows and their history. The caller commits."""
    today = date.today()
    count = 0
    for r in rows:
        status = status_map.get(r["status"]) or None
        cow = Cow(
            ear_tag=r["tag"],
            date_of_birth=date.fromisoformat(r["birth"]) if r["birth"] else None,
            gender=Cow.GENDER_MALE if status == Cow.BS_B_CALF else Cow.GENDER_FEMALE,
            group_id=group_map[r["group"]],
            status=Cow.STATUS_ACTIVE,
            breeding_status=status,
            season_count_offset=r["lactation"],
            notes="استيراد من برنامج المحلب",
            created_by_id=user_id,
        )
        db.session.add(cow)
        db.session.flush()
        db.session.add(CowMovement(
            cow_id=cow.id, from_group_id=None, to_group_id=cow.group_id,
            moved_on=today, reason="استيراد من برنامج المحلب",
            created_by_id=user_id,
        ))
        if status:
            db.session.add(BreedingStatusChange(
                cow_id=cow.id, from_status=None, to_status=status,
                changed_by_id=user_id, reason="استيراد من برنامج المحلب",
            ))
        if r["last_calving"]:
            db.session.add(BreedingEvent(
                cow_id=cow.id, event_type=BreedingEvent.EVENT_CALVING,
                event_date=date.fromisoformat(r["last_calving"]),
                notes="آخر ولادة — استيراد", created_by_id=user_id,
            ))
        if r["last_insem"]:
            db.session.add(BreedingEvent(
                cow_id=cow.id, event_type=BreedingEvent.EVENT_INSEMINATION,
                event_date=date.fromisoformat(r["last_insem"]),
                notes="آخر تلقيح — استيراد", created_by_id=user_id,
            ))
        count += 1
    return count
