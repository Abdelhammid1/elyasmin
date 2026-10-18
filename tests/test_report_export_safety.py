"""Export helpers hardening (security review of HERD-6).

  1. PDF render URL ignores the request's host — a spoofed Host header
     can't point server-side Chromium at another server.
  2. Excel export neutralises text that would run as a formula.
"""
from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from app.utils.reports import excel_response, local_render_url


def test_pdf_url_uses_fixed_base(app):
    with app.test_request_context("/", base_url="http://evil.example:80"):
        url, domain = local_render_url(
            "http://evil.example/herd/?status=preg&group=1&print=1")
        base = app.config["PDF_RENDER_BASE_URL"].rstrip("/")
        assert url == base + "/herd/?status=preg&group=1&print=1"
        assert "evil.example" not in url
        assert domain == "127.0.0.1"


def test_pdf_url_without_query(app):
    with app.test_request_context("/"):
        url, _ = local_render_url("http://anything/finance/pnl")
        assert url.endswith("/finance/pnl")


def test_excel_formula_injection_neutralised(app):
    with app.test_request_context("/"):
        resp = excel_response(
            "x", ["h"],
            [["=HYPERLINK(\"http://x\",\"y\")"], ["+1+2"], ["@SUM(A1)"],
             ["-3"], ["102"], [42]],
            "x.xlsx",
        )
    ws = load_workbook(BytesIO(resp.get_data())).active
    vals = [row[0] for row in ws.iter_rows(min_row=2, values_only=True)]
    assert vals[0] == "'=HYPERLINK(\"http://x\",\"y\")"
    assert vals[1] == "'+1+2"
    assert vals[2] == "'@SUM(A1)"
    assert vals[3] == "'-3"
    assert vals[4] == "102"
    assert vals[5] == 42
    # Nothing was stored as a formula.
    assert all(c.data_type != "f" for row in ws.iter_rows() for c in row)
