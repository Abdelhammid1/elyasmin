"""HERD-7 regression suite — herd alerts.

  1. Pregnancy check: insem 41 days ago → listed; 39 days → not.
  2. Drying: preg 211 days → listed; status dry → not in drying.
  3. Waiting: dry with calving in 15 days → listed; status waiting → not.
  4. Recording the pregnancy check removes the cow from the list.
  5. Settings change the threshold.
  6. Excel export of one list; event form preselects ?type=.
"""
from __future__ import annotations

import re
from datetime import date, timedelta
from io import BytesIO

from openpyxl import load_workbook

from app.extensions import db
from app.models.finance import Setting
from app.models.herd import (
    BreedingEvent, BreedingStatusChange, CattleGroup, Cow, CowMovement,
)
from app.services.herd_alerts import build_alerts

_PREFIX = "TEST-HERD7-"


def _make(app, tag, status, insem_days_ago):
    with app.app_context():
        c = Cow(ear_tag=_PREFIX + tag, gender=Cow.GENDER_FEMALE,
                group_id=CattleGroup.query.filter_by(type="milk", is_archived=False).first().id,
                status=Cow.STATUS_ACTIVE, breeding_status=status)
        db.session.add(c)
        db.session.flush()
        db.session.add(BreedingEvent(
            cow_id=c.id, event_type=BreedingEvent.EVENT_INSEMINATION,
            event_date=date.today() - timedelta(days=insem_days_ago)))
        db.session.commit()
        return c.id


def _cleanup(app):
    with app.app_context():
        for c in Cow.query.filter(Cow.ear_tag.like(_PREFIX + "%")).all():
            BreedingStatusChange.query.filter_by(cow_id=c.id).delete()
            BreedingEvent.query.filter_by(cow_id=c.id).delete()
            CowMovement.query.filter_by(cow_id=c.id).delete()
            db.session.delete(c)
        for key in (Setting.KEY_HERD_ALERT_CHECK_DAYS,):
            row = db.session.get(Setting, key)
            if row:
                db.session.delete(row)
        db.session.commit()


def _tags(app, rule):
    with app.app_context():
        return [r["tag"] for r in build_alerts()[rule] if r["tag"].startswith(_PREFIX)]


def test_check_rule(app):
    _cleanup(app)
    try:
        _make(app, "C41", Cow.BS_INSEM, 41)
        _make(app, "C39", Cow.BS_INSEM, 39)
        assert _tags(app, "check") == [_PREFIX + "C41"]
    finally:
        _cleanup(app)


def test_dry_rule(app):
    _cleanup(app)
    try:
        _make(app, "D211", Cow.BS_PREG, 211)
        _make(app, "D200", Cow.BS_PREG2, 200)
        _make(app, "DDRY", Cow.BS_DRY, 230)
        assert _tags(app, "dry") == [_PREFIX + "D211"]
    finally:
        _cleanup(app)


def test_waiting_rule(app):
    _cleanup(app)
    try:
        _make(app, "W265", Cow.BS_DRY, 265)    # due in 15 days
        _make(app, "W250", Cow.BS_DRY, 250)    # due in 30 days
        _make(app, "WWAIT", Cow.BS_WAITING, 270)
        assert _tags(app, "waiting") == [_PREFIX + "W265"]
    finally:
        _cleanup(app)


def test_recording_check_clears_alert(admin_client, app):
    _cleanup(app)
    try:
        cid = _make(app, "R45", Cow.BS_INSEM, 45)
        assert _tags(app, "check") == [_PREFIX + "R45"]
        admin_client.post(f"/herd/{cid}/breeding/new", data={
            "event_type": "pregnancy_check",
            "event_date": date.today().isoformat(),
            "result": "pregnant",
        })
        assert _tags(app, "check") == []
    finally:
        _cleanup(app)


def test_settings_change_threshold(admin_client, app):
    _cleanup(app)
    try:
        _make(app, "S35", Cow.BS_INSEM, 35)
        assert _tags(app, "check") == []
        r = admin_client.post("/herd/alert-settings", data={
            Setting.KEY_HERD_ALERT_CHECK_DAYS: "30",
            Setting.KEY_HERD_ALERT_DRY_DAYS: "210",
            Setting.KEY_HERD_ALERT_WAITING_DAYS: "20",
            Setting.KEY_GESTATION_DAYS: "280",
        })
        assert r.status_code in (302, 303)
        assert _tags(app, "check") == [_PREFIX + "S35"]
    finally:
        _cleanup(app)


def test_excel_and_event_preselect(admin_client, app):
    _cleanup(app)
    try:
        cid = _make(app, "X50", Cow.BS_INSEM, 50)
        r = admin_client.get("/herd/alerts.xlsx?rule=check")
        assert r.status_code == 200
        ws = load_workbook(BytesIO(r.data)).active
        tags = [row[0] for row in ws.iter_rows(min_row=2, values_only=True)]
        assert _PREFIX + "X50" in tags

        body = admin_client.get(f"/herd/{cid}/breeding/new?type=drying").get_data(as_text=True)
        opt = re.search(r'<option[^>]*value="drying"[^>]*>', body).group(0)
        assert "selected" in opt, opt

        assert admin_client.get("/herd/alerts").status_code == 200
        assert admin_client.get("/").status_code in (200, 302)
    finally:
        _cleanup(app)
