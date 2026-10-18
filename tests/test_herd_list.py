"""HERD-6 regression suite — herd list columns, filters, export.

  1. Computed columns (season, DIM, days pregnant, expected calving).
  2. Combined status + group filter; status=none.
  3. Season filter.
  4. Sort by days pregnant, descending.
  5. Excel export holds exactly the filtered cows.
  6. PDF export route registered.
"""
from __future__ import annotations

from datetime import date, timedelta
from io import BytesIO

from openpyxl import load_workbook

from app.extensions import db
from app.models.herd import Birth, BreedingEvent, CattleGroup, Cow
from app.services.herd_list import build_herd_rows

_PREFIX = "TEST-HERD6-"


def _gid(app, gtype):
    with app.app_context():
        return CattleGroup.query.filter_by(type=gtype, is_archived=False).first().id


def _make(app, tag, gtype="milk", status=None, offset=0,
          calved_days_ago=None, insem_days_ago=None):
    with app.app_context():
        today = date.today()
        c = Cow(ear_tag=_PREFIX + tag, gender=Cow.GENDER_FEMALE,
                group_id=CattleGroup.query.filter_by(type=gtype, is_archived=False).first().id,
                status=Cow.STATUS_ACTIVE, breeding_status=status,
                season_count_offset=offset,
                date_of_birth=today - timedelta(days=4 * 365))
        db.session.add(c)
        db.session.flush()
        if calved_days_ago is not None:
            db.session.add(Birth(mother_id=c.id, calves_count=1,
                                 birth_date=today - timedelta(days=calved_days_ago),
                                 delivery_type="natural"))
        if insem_days_ago is not None:
            db.session.add(BreedingEvent(
                cow_id=c.id, event_type=BreedingEvent.EVENT_INSEMINATION,
                event_date=today - timedelta(days=insem_days_ago)))
        db.session.commit()
        return c.id


def _cleanup(app):
    with app.app_context():
        for c in Cow.query.filter(Cow.ear_tag.like(_PREFIX + "%")).all():
            BreedingEvent.query.filter_by(cow_id=c.id).delete()
            Birth.query.filter_by(mother_id=c.id).delete()
            db.session.delete(c)
        db.session.commit()


def _mine(rows):
    return [r for r in rows if r["tag"].startswith(_PREFIX)]


def test_computed_columns(app):
    _cleanup(app)
    try:
        _make(app, "A", status="preg", offset=1,
              calved_days_ago=100, insem_days_ago=50)
        with app.app_context():
            r = _mine(build_herd_rows({"q": _PREFIX, "life": "active"}))[0]
            today = date.today()
            assert r["season"] == 2
            assert r["dim"] == 100
            assert r["days_preg"] == 50
            assert r["expected"] == today - timedelta(days=50) + timedelta(days=280)
            assert r["status"] == "عشار"
    finally:
        _cleanup(app)


def test_no_days_pregnant_when_not_pregnant(app):
    _cleanup(app)
    try:
        _make(app, "B", status="insem", calved_days_ago=60, insem_days_ago=10)
        with app.app_context():
            r = _mine(build_herd_rows({"q": _PREFIX, "life": "active"}))[0]
            assert r["days_preg"] is None and r["expected"] is None
            assert r["last_insem"] == date.today() - timedelta(days=10)
    finally:
        _cleanup(app)


def test_combined_filters(admin_client, app):
    _cleanup(app)
    milk, dry = _gid(app, "milk"), _gid(app, "dry")
    try:
        _make(app, "P1", "milk", status="preg")
        _make(app, "P2", "dry", status="preg")
        _make(app, "I1", "milk", status="insem")
        _make(app, "N1", "milk", status=None)
        with app.app_context():
            tags = [r["tag"] for r in _mine(build_herd_rows(
                {"q": _PREFIX, "status": ["preg"], "group": [milk]}))]
            assert tags == [_PREFIX + "P1"]
            tags = [r["tag"] for r in _mine(build_herd_rows(
                {"q": _PREFIX, "status": ["none"]}))]
            assert tags == [_PREFIX + "N1"]
        body = admin_client.get(
            f"/herd/?q={_PREFIX}&status=preg&group={milk}&group={dry}"
        ).get_data(as_text=True)
        assert _PREFIX + "P1" in body and _PREFIX + "P2" in body
        assert _PREFIX + "I1" not in body
    finally:
        _cleanup(app)


def test_season_filter(app):
    _cleanup(app)
    try:
        _make(app, "S0", offset=0)
        _make(app, "S2", offset=2)
        with app.app_context():
            tags = [r["tag"] for r in _mine(build_herd_rows(
                {"q": _PREFIX, "season": [2]}))]
            assert tags == [_PREFIX + "S2"]
    finally:
        _cleanup(app)


def test_sort_days_pregnant_desc(app):
    _cleanup(app)
    try:
        _make(app, "D1", status="preg", insem_days_ago=30)
        _make(app, "D2", status="preg", insem_days_ago=90)
        _make(app, "D3", status=None)
        with app.app_context():
            tags = [r["tag"] for r in _mine(build_herd_rows(
                {"q": _PREFIX}, sort="days_preg", direction="desc"))]
            assert tags == [_PREFIX + "D2", _PREFIX + "D1", _PREFIX + "D3"]
    finally:
        _cleanup(app)


def test_excel_export_matches_filter(admin_client, app):
    _cleanup(app)
    try:
        _make(app, "E1", status="preg")
        _make(app, "E2", status="dry")
        r = admin_client.get(f"/herd/export.xlsx?q={_PREFIX}&status=preg")
        assert r.status_code == 200
        ws = load_workbook(BytesIO(r.data)).active
        rows = list(ws.iter_rows(values_only=True))
        assert rows[0][0] == "رقم الأذن"
        assert [row[0] for row in rows[1:]] == [_PREFIX + "E1"]
    finally:
        _cleanup(app)


def test_pdf_route_registered(app):
    rules = [r for r in app.url_map.iter_rules() if r.endpoint == "herd.export_cows_pdf"]
    assert len(rules) == 1
