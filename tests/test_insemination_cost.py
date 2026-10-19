"""HERD-8 regression suite — insemination cost.

  1. From stock: draws the straws at average cost, StockMovement,
     JE DR 5025 / CR inventory.
  2. Not enough stock → refused, nothing saved.
  3. Cash: Expense (insemination), treasury down, JE DR 5025 / CR treasury.
  4. No cost → event only.
  5. Cost report + Excel.
"""
from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from app.extensions import db
from app.models.accounting import JournalEntry, JournalLine
from app.models.finance import AccountMovement, Expense, TreasuryAccount
from app.models.herd import (
    BreedingEvent, BreedingStatusChange, CattleGroup, Cow, CowMovement,
)
from app.models.inventory import Ingredient, StockMovement

_TAG = "TEST-HERD8-COW"
_ITEM = "TEST-HERD8-SEMEN"


def _seed(app, qty=10, avg=50):
    with app.app_context():
        c = Cow.query.filter_by(ear_tag=_TAG).first()
        if c is None:
            c = Cow(ear_tag=_TAG, gender=Cow.GENDER_FEMALE,
                    group_id=CattleGroup.query.filter_by(type="milk").first().id,
                    status=Cow.STATUS_ACTIVE, breeding_status=Cow.BS_OPEN)
            db.session.add(c)
        ing = Ingredient.query.filter_by(name=_ITEM).first()
        if ing is None:
            ing = Ingredient(name=_ITEM, category="سائل منوي", unit=Ingredient.UNIT_PIECE)
            db.session.add(ing)
        ing.current_qty = Decimal(qty)
        ing.avg_cost = Decimal(avg)
        db.session.commit()
        return c.id, ing.id


def _treasury(app):
    with app.app_context():
        t = TreasuryAccount.query.filter_by(is_archived=False).first()
        return t.id, Decimal(t.current_balance or 0)


def _cleanup(app):
    with app.app_context():
        c = Cow.query.filter_by(ear_tag=_TAG).first()
        if c is not None:
            for e in BreedingEvent.query.filter_by(cow_id=c.id).all():
                for je in JournalEntry.query.filter_by(source_type="BreedingEvent", source_id=e.id).all():
                    JournalLine.query.filter_by(entry_id=je.id).delete()
                    db.session.delete(je)
                StockMovement.query.filter_by(reason=StockMovement.REASON_INSEMINATION,
                                              ref_id=e.id).delete()
                if e.cost_expense_id:
                    exp = db.session.get(Expense, e.cost_expense_id)
                    for je in JournalEntry.query.filter_by(source_type="Expense", source_id=exp.id).all():
                        JournalLine.query.filter_by(entry_id=je.id).delete()
                        db.session.delete(je)
                    for mv in AccountMovement.query.filter_by(ref_type="expense", ref_id=exp.id).all():
                        mv.account.current_balance = (mv.account.current_balance or 0) - mv.amount
                        db.session.delete(mv)
                    e.cost_expense_id = None
                    db.session.flush()
                    db.session.delete(exp)
            BreedingStatusChange.query.filter_by(cow_id=c.id).delete()
            BreedingEvent.query.filter_by(cow_id=c.id).delete()
            CowMovement.query.filter_by(cow_id=c.id).delete()
            db.session.delete(c)
        Ingredient.query.filter_by(name=_ITEM).delete()
        db.session.commit()


def _post(admin_client, cow_id, **cost):
    data = {"event_type": "insemination", "event_date": date.today().isoformat()}
    data.update({k: str(v) for k, v in cost.items()})
    return admin_client.post(f"/herd/{cow_id}/breeding/new", data=data,
                             follow_redirects=False)


def _je_lines(app, source_type, source_id):
    with app.app_context():
        # Newest first: other tests can leave an orphan JE behind for a
        # deleted Expense, and SQLite reuses the id.
        je = (JournalEntry.query.filter_by(source_type=source_type, source_id=source_id)
              .order_by(JournalEntry.id.desc()).first())
        if je is None:
            return {}
        return {(l.account.code if l.account.code in ("5025", "1200") else "other"):
                (l.debit, l.credit) for l in je.lines}


def test_cost_from_stock(admin_client, app):
    _cleanup(app)
    cow_id, ing_id = _seed(app)
    try:
        r = _post(admin_client, cow_id, cost_source="stock",
                  cost_ingredient_id=ing_id, cost_qty=1)
        assert r.status_code in (302, 303), r.data[:300]
        with app.app_context():
            e = BreedingEvent.query.filter_by(cow_id=cow_id).first()
            assert e.cost_source == "stock" and e.cost_amount == Decimal("50.00")
            assert db.session.get(Ingredient, ing_id).current_qty == Decimal("9")
            mv = StockMovement.query.filter_by(ref_id=e.id,
                                               reason=StockMovement.REASON_INSEMINATION).first()
            assert mv.delta == Decimal("-1")
            eid = e.id
        lines = _je_lines(app, "BreedingEvent", eid)
        assert lines["5025"] == (Decimal("50.00"), Decimal("0"))
        assert lines["1200"] == (Decimal("0"), Decimal("50.00"))
    finally:
        _cleanup(app)


def test_not_enough_stock_refused(admin_client, app):
    _cleanup(app)
    cow_id, ing_id = _seed(app, qty=1)
    try:
        r = _post(admin_client, cow_id, cost_source="stock",
                  cost_ingredient_id=ing_id, cost_qty=3)
        assert r.status_code == 200  # form re-rendered with the error
        with app.app_context():
            assert BreedingEvent.query.filter_by(cow_id=cow_id).count() == 0
            assert db.session.get(Ingredient, ing_id).current_qty == Decimal("1")
    finally:
        _cleanup(app)


def test_cost_cash(admin_client, app):
    _cleanup(app)
    cow_id, _ = _seed(app)
    tid, before = _treasury(app)
    try:
        r = _post(admin_client, cow_id, cost_source="cash",
                  cost_amount="300", cost_account_id=tid)
        assert r.status_code in (302, 303), r.data[:300]
        with app.app_context():
            e = BreedingEvent.query.filter_by(cow_id=cow_id).first()
            assert e.cost_source == "cash" and e.cost_amount == Decimal("300.00")
            exp = db.session.get(Expense, e.cost_expense_id)
            assert exp.category == Expense.CAT_INSEMINATION
            assert Decimal(db.session.get(TreasuryAccount, tid).current_balance) == before - 300
            exp_id = exp.id
        lines = _je_lines(app, "Expense", exp_id)
        assert lines["5025"] == (Decimal("300.00"), Decimal("0"))
    finally:
        _cleanup(app)


def test_no_cost(admin_client, app):
    _cleanup(app)
    cow_id, _ = _seed(app)
    try:
        _post(admin_client, cow_id, cost_source="")
        with app.app_context():
            e = BreedingEvent.query.filter_by(cow_id=cow_id).first()
            assert e.cost_source is None and e.cost_amount is None
            assert JournalEntry.query.filter_by(source_type="BreedingEvent",
                                                source_id=e.id).count() == 0
    finally:
        _cleanup(app)


def test_report_and_excel(admin_client, app):
    _cleanup(app)
    cow_id, ing_id = _seed(app)
    try:
        _post(admin_client, cow_id, cost_source="stock",
              cost_ingredient_id=ing_id, cost_qty=2)
        today = date.today().isoformat()
        body = admin_client.get(f"/herd/insemination-costs?from={today}&to={today}").get_data(as_text=True)
        assert _TAG in body and "100.00" in body
        r = admin_client.get(f"/herd/insemination-costs.xlsx?from={today}&to={today}")
        rows = list(load_workbook(BytesIO(r.data)).active.iter_rows(min_row=2, values_only=True))
        assert any(row[1] == _TAG and row[5] == 100 for row in rows)
    finally:
        _cleanup(app)


def test_form_defaults_to_stock(admin_client, app):
    """Decided default: with semen in stock the form opens on "from stock"."""
    import re
    _cleanup(app)
    cow_id, ing_id = _seed(app)
    try:
        body = admin_client.get(f"/herd/{cow_id}/breeding/new").get_data(as_text=True)
        assert re.search(r'<option selected value="stock">', body)
        assert re.search(rf'<option selected value="{ing_id}">', body) or \
            f'value="{ing_id}"' in body   # several semen items → none pre-picked
    finally:
        _cleanup(app)
